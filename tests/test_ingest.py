"""Phase 1: loaders take real quirks in their stride and refuse anything that changed shape."""
from __future__ import annotations

import gzip
import json
from datetime import date

import openpyxl
import pandas as pd
import pytest

from src import config, ingest_engagement, ingest_imdb, ingest_top10, ingest_wikipedia, manifest
from src.schema import SchemaError, expect_not_html

PERIOD = config.Period("2026H1", date(2026, 1, 1), date(2026, 6, 30), "unused")
HEADER = ["Title", "Available Globally?", "Release Date", "Hours Viewed", "Runtime", "Views"]


# --- engagement report ---------------------------------------------------------------

def write_report(path, sheets):
    """Same layout as the real files: five preamble rows, header on row 6, data from column B."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, (header, rows) in sheets.items():
        ws = wb.create_sheet(name)
        ws["B1"] = "Netflix's What We Watched Report"
        ws["B2"] = "Views from January to June 2026* (see footnote)"
        for i, v in enumerate(header):
            ws.cell(row=6, column=2 + i, value=v)
        for r, row in enumerate(rows, start=7):
            for i, v in enumerate(row):
                ws.cell(row=r, column=2 + i, value=v)
    wb.save(path)
    return path


TV_ROWS = [
    ["His & Hers: Limited Series", "Yes", "2026-01-08", 454300000, "4:22", 104000000],
    ["Peppa Pig: Season 1", "No", None, 50000000, "*", 9000000],                 # runtime withheld
    ["Raw: January 5, 2026", "Yes", "2026-01-05", 3000000, "2:59*", 1000000],    # starred runtime
    ["Valle Salvaje: Season 2 // Valle Salvaje: Temporada 2", "No", None, 9000000, "184:10", 50000],
    ["Other Shows", "*", "*", 757000000, "*", "*"],
    [],
    ["*“Other Movies” and “Other Shows” rows in this report capture total viewing"],
]
FILM_ROWS = [["War Machine", "Yes", "2026-03-06", 266800000, "1:49", 146900000],
             [1899, "No", None, 100000, "1:30", 60000]]


@pytest.fixture
def report(tmp_path):
    return write_report(tmp_path / "r.xlsx", {"Shows": (HEADER, TV_ROWS), "Movies": (HEADER, FILM_ROWS)})


def test_engagement_loads_titles_verbatim_with_quirks_flagged(report):
    df = ingest_engagement.read_workbook(report, PERIOD).set_index("title")
    assert len(df) == 7
    assert df.loc["Valle Salvaje: Season 2 // Valle Salvaje: Temporada 2", "runtime_hours"] == pytest.approx(184 + 10 / 60)
    peppa = df.loc["Peppa Pig: Season 1"]
    assert peppa.runtime_withheld and pd.isna(peppa.runtime_hours)
    assert peppa.views == 9000000
    raw = df.loc["Raw: January 5, 2026"]
    assert raw.runtime_starred and not raw.runtime_withheld and raw.runtime_hours == pytest.approx(2 + 59 / 60)
    assert df.loc["1899", "kind"] == "film"  # a numeric title is still a title


def test_catch_all_rows_are_flagged_not_parsed(report):
    df = ingest_engagement.read_workbook(report, PERIOD)
    other = df[df.is_catch_all]
    assert other.title.tolist() == ["Other Shows"]
    assert other.hours_viewed.item() == 757000000
    assert other.views.isna().all() and not other.runtime_withheld.item()


@pytest.mark.parametrize("header", [
    HEADER[:5],                                                     # Views dropped
    ["Title", "Available Globally?", "Release Date", "Hours", "Runtime", "Views"],  # renamed
    ["Title", "Release Date", "Available Globally?", "Hours Viewed", "Runtime", "Views"],  # reordered
])
def test_changed_columns_are_refused(tmp_path, header):
    path = write_report(tmp_path / "r.xlsx", {"Shows": (header, TV_ROWS[:1]), "Movies": (HEADER, FILM_ROWS)})
    with pytest.raises(SchemaError, match="columns changed"):
        ingest_engagement.read_workbook(path, PERIOD)


def test_unexpected_sheets_are_refused(tmp_path):
    path = write_report(tmp_path / "r.xlsx", {"Engagement": (HEADER, TV_ROWS[:1])})
    with pytest.raises(SchemaError, match="one TV and one film sheet"):
        ingest_engagement.read_workbook(path, PERIOD)


@pytest.mark.parametrize("row, message", [
    (["X", "Yes", None, 100000, "1h30", 5], "runtime"),
    (["X", "Yes", None, 100000, "1:30", "*"], "views"),     # '*' is only accepted on catch-all rows
    (["X", "*", None, 100000, "1:30", 5], "Available Globally"),
    (["X", "Yes", None, "lots", "1:30", 5], "hours"),
])
def test_unexpected_values_are_refused(tmp_path, row, message):
    path = write_report(tmp_path / "r.xlsx", {"TV": (HEADER, [row]), "Film": (HEADER, FILM_ROWS)})
    with pytest.raises(SchemaError, match=message):
        ingest_engagement.read_workbook(path, PERIOD)


# --- Top 10 --------------------------------------------------------------------------

CATEGORIES = ["Films (English)", "Films (Non-English)", "TV (English)", "TV (Non-English)"]


def top10_text(weeks=("2023-06-11", "2023-06-18"), drop_last=False, header=config.TOP10_COLUMNS):
    lines = ["\t".join(header)]
    for week in weeks:
        views = "N/A" if week < "2023-06-18" else "1000000"
        runtime = "N/A" if views == "N/A" else "1.5"
        for cat in CATEGORIES:
            for rank in range(1, 11):
                title, season = ("NA", "N/A") if cat.startswith("Films") else ("Show", "Show: Season 1")
                lines.append("\t".join([week, cat, str(rank), title, season, "2000000", runtime, views, "1"]))
    if drop_last:
        lines.pop()
    return "\n".join(lines) + "\n"


def test_top10_parses_na_as_missing_but_keeps_a_title_called_na():
    df = ingest_top10.read_tsv(top10_text())
    assert len(df) == 80
    assert (df[df.category.str.startswith("Films")].show_title == "NA").all()
    assert df[df.category.str.startswith("Films")].season_title.isna().all()
    early = df[df.week == date(2023, 6, 11)]
    assert early.weekly_views.isna().all() and early.runtime.isna().all()
    assert df[df.week == date(2023, 6, 18)].weekly_views.notna().all()
    assert set(df.kind) == {"tv", "film"} and set(df.language) == {"english", "non-english"}


def test_top10_refuses_changed_header_and_short_weeks():
    with pytest.raises(SchemaError, match="columns changed"):
        ingest_top10.read_tsv(top10_text(header=config.TOP10_COLUMNS[:-1] + ["weeks_in_top_10"]))
    with pytest.raises(SchemaError, match="10 rows per week"):
        ingest_top10.read_tsv(top10_text(drop_last=True))


class FakeResponse:
    def __init__(self, content, ctype):
        self.content, self.headers = content, {"Content-Type": ctype}


def test_an_unsupported_browser_page_is_not_saved_as_data():
    with pytest.raises(SchemaError, match="HTML"):
        expect_not_html(FakeResponse(b"<!doctype html><html>", "text/html; charset=utf-8"), "u")
    with pytest.raises(SchemaError, match="HTML"):  # mislabelled content type, same page
        expect_not_html(FakeResponse(b"<!doctype html><html>", "text/tab-separated-values"), "u")
    expect_not_html(FakeResponse(b"week\tcategory\n", "text/tab-separated-values"), "u")


# --- manifest ------------------------------------------------------------------------

def test_manifest_catches_a_replaced_raw_file(tmp_path):
    raw, log = tmp_path / "file.tsv", tmp_path / "manifest.json"
    raw.write_text("a\tb\n1\t2\n")
    entry = manifest.record(raw, "test", "https://example.org/file.tsv", manifest=log)
    assert entry["bytes"] == raw.stat().st_size and len(entry["sha256"]) == 64
    assert manifest.verify(raw, manifest=log)["url"] == "https://example.org/file.tsv"
    raw.write_text("a\tb\n1\t3\n")
    with pytest.raises(ValueError, match="no longer matches"):
        manifest.verify(raw, manifest=log)
    with pytest.raises(FileNotFoundError):
        manifest.verify(tmp_path / "never-downloaded.tsv", manifest=log)


# --- IMDb ----------------------------------------------------------------------------

def write_gz(path, lines):
    path.parent.mkdir(parents=True, exist_ok=True)
    with gzip.open(path, "wt", encoding="utf-8", newline="\n") as f:
        f.write("\n".join(lines) + "\n")


def test_imdb_loads_bare_quotes_and_null_markers_without_merging_rows(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "RAW", tmp_path)
    write_gz(ingest_imdb.raw_path("title.basics"), [
        "\t".join(config.IMDB_COLUMNS["title.basics"]),
        'tt0000001\tmovie\tThe "Real" One\tThe "Real" One\t0\t1999\t\\N\t101\tDrama',
        'tt0000002\ttvSeries\t"Quoted\tQuoted\t0\t2020\t2022\t\\N\tComedy,Drama',  # unbalanced quote
        "tt0000003\ttvEpisode\tPilot\tPilot\t0\t\\N\t\\N\t\\N\t\\N",
    ])
    con = ingest_imdb.build(names=("title.basics",), db_path=tmp_path / "imdb.duckdb", check_manifest=False)
    rows = con.execute("SELECT tconst, primaryTitle, endYear, runtimeMinutes, genres FROM title_basics ORDER BY 1").fetchall()
    assert rows == [("tt0000001", 'The "Real" One', None, 101, "Drama"),
                    ("tt0000002", '"Quoted', 2022, None, "Comedy,Drama"),
                    ("tt0000003", "Pilot", None, None, None)]


def test_imdb_refuses_changed_header(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "RAW", tmp_path)
    write_gz(ingest_imdb.raw_path("title.ratings"), ["tconst\taverageRating\tvotes", "tt0000001\t7.1\t10"])
    with pytest.raises(SchemaError, match="columns changed"):
        ingest_imdb.build(names=("title.ratings",), db_path=tmp_path / "imdb.duckdb", check_manifest=False)


# --- Wikipedia -----------------------------------------------------------------------

def test_wikidata_mapping_keeps_misses_and_one_to_many(tmp_path, monkeypatch):
    monkeypatch.setattr(manifest, "verify", lambda path: {})
    item = lambda q: {"type": "uri", "value": f"http://www.wikidata.org/entity/{q}"}
    art = lambda a: {"type": "uri", "value": f"https://en.wikipedia.org/wiki/{a}"}
    lit = lambda v: {"type": "literal", "value": v}
    payload = {"imdb_ids": ["tt1", "tt2", "tt3"], "response": {"results": {"bindings": [
        {"imdb": lit("tt1"), "item": item("Q1"), "article": art("Wednesday_(TV_series)")},
        {"imdb": lit("tt2"), "item": item("Q2")},                                  # item, no English article
        {"imdb": lit("tt2"), "item": item("Q3"), "article": art("Duplicate")},     # two items share an id
    ]}}}
    path = tmp_path / "batch.json"
    path.write_text(json.dumps(payload))
    df = ingest_wikipedia.load_articles([path])
    assert len(df) == 4
    assert df[df.imdb_id == "tt2"].wikidata_item.tolist() == ["Q2", "Q3"]
    assert df[df.imdb_id == "tt3"].wikidata_item.isna().all()  # no Wikidata item at all: kept as a miss


def test_pageviews_loads_daily_user_views_and_refuses_anything_else(tmp_path, monkeypatch):
    monkeypatch.setattr(manifest, "verify", lambda path: {})
    item = {"project": "en.wikipedia", "article": "X", "granularity": "daily",
            "timestamp": "2026091500", "access": "all-access", "agent": "user", "views": 3203}
    good = tmp_path / "good.json"
    good.write_text(json.dumps({"article": "X", "response": {"items": [item]}}))
    empty = tmp_path / "empty.json"
    empty.write_text(json.dumps({"article": "Y", "response": {"items": [], "status": 404}}))
    df = ingest_wikipedia.load_pageviews([good, empty])
    assert df.to_dict("records") == [{"article": "X", "date": date(2026, 9, 15), "views": 3203}]
    bad = tmp_path / "bad.json"
    bad.write_text(json.dumps({"article": "X", "response": {"items": [dict(item, granularity="monthly")]}}))
    with pytest.raises(SchemaError):
        ingest_wikipedia.load_pageviews([bad])


def test_wikidata_query_rejects_non_imdb_ids():
    with pytest.raises(ValueError, match="not IMDb title ids"):
        ingest_wikipedia.fetch_articles(["tt13443470", "Wednesday"])
